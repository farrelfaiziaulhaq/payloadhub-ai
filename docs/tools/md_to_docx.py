"""Convert the project Markdown reports to DOCX without external tools.

Usage:
    python docs/tools/md_to_docx.py

The converter intentionally supports only the Markdown subset used by the
reports in this repository: headings, paragraphs, fenced code, tables,
blockquotes, ordered/unordered lists, horizontal rules, and inline bold,
italic, code, and links. It is deterministic and does not touch the network.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, List, Optional, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCS_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = DOCS_DIR / "docx"

SOURCES = [
    "Tahap-1-Analisis-Kebutuhan.md",
    "Tahap-2-Perancangan-Sistem.md",
    "Tahap-3-Implementasi.md",
    "Tahap-4-Pengujian.md",
    "Tahap-5-Laporan-Akhir.md",
    "DOKUMENTASI-TEKNIS.md",
    "PANDUAN-DEMO-VIDEO.md",
    "BAHAN-PRESENTASI.md",
    "README.md",
]

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)",
)
TABLE_SEPARATOR_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
ORDERED_RE = re.compile(r"^\s*(\d+)[.)]\s+(.*)$")
BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")


def set_cell_background(cell, hex_color: str) -> None:
    """Apply a solid background color to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_inline(paragraph, text: str, base_bold: bool = False) -> None:
    """Add text with **bold**, *italic*, `code`, and [link](url) formatting."""
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            run.bold = base_bold
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x9C, 0x27, 0x0B)
        elif token.startswith("[") and "](" in token:
            label, url = token[1:-1].split("](", 1)
            run = paragraph.add_run(label)
            run.underline = True
            run.font.color.rgb = RGBColor(0x1A, 0x4F, 0x8B)
            if url and not url.startswith("#"):
                # Plain URL suffix keeps the target visible in printed reports;
                # the label itself remains the readable link text.
                suffix = paragraph.add_run(f" ({url})")
                suffix.font.size = Pt(8.5)
                suffix.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold = base_bold


def add_code_block(document: Document, lines: List[str]) -> None:
    if not lines:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.left_indent = Inches(0.2)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "F4F4F4")
    paragraph._p.get_or_add_pPr().append(shading)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)


def split_row(line: str) -> List[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def add_table(document: Document, rows: List[List[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index in range(width):
            value = row[column_index] if column_index < len(row) else ""
            paragraph = cells[column_index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, base_bold=(row_index == 0))
            if row_index == 0:
                set_cell_background(cells[column_index], "E8EDF3")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9.5)
            else:
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
    document.add_paragraph()


def convert(markdown_path: Path, output_path: Path) -> None:
    text = markdown_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)

    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0x14, 0x2B, 0x45)

    index = 0
    in_code = False
    code_lines: List[str] = []
    table_lines: List[str] = []
    first_heading = True

    def flush_table() -> None:
        nonlocal table_lines
        if not table_lines:
            return
        rows = [
            split_row(line)
            for line in table_lines
            if not TABLE_SEPARATOR_RE.match(line.strip())
        ]
        add_table(document, rows)
        table_lines = []

    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if line.strip().startswith("|") and "|" in line:
            table_lines.append(line.strip())
            index += 1
            continue
        flush_table()

        if not line.strip():
            index += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            content = heading.group(2).strip()
            if first_heading and level == 1:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = paragraph.add_run(content)
                run.bold = True
                run.font.size = Pt(20)
                run.font.color.rgb = RGBColor(0x14, 0x2B, 0x45)
                paragraph.paragraph_format.space_after = Pt(8)
                first_heading = False
            else:
                document.add_heading("", level=min(level, 4))
                paragraph = document.paragraphs[-1]
                add_inline(paragraph, content)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0x14, 0x2B, 0x45)
            index += 1
            continue

        if line.strip() in {"---", "***", "___"}:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run("─" * 60)
            run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
            run.font.size = Pt(8)
            index += 1
            continue

        if line.lstrip().startswith(">"):
            content = line.lstrip()[1:].strip()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, content)
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            index += 1
            continue

        ordered = ORDERED_RE.match(line)
        if ordered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, ordered.group(2))
            index += 1
            continue

        bullet = BULLET_RE.match(line)
        if bullet:
            checkbox = re.match(r"^\[( |x|X)\]\s+(.*)$", bullet.group(1))
            if checkbox:
                marker = "☑" if checkbox.group(1).lower() == "x" else "☐"
                paragraph = document.add_paragraph()
                add_inline(paragraph, f"{marker} {checkbox.group(2)}")
            else:
                paragraph = document.add_paragraph(style="List Bullet")
                add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, line.strip())
        index += 1

    if in_code:
        add_code_block(document, code_lines)
    flush_table()

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("PayloadHub-AI — Farrel Faizi Aulhaq")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    converted = 0
    for name in SOURCES:
        source = DOCS_DIR / name
        if not source.exists():
            print(f"[skip] {name} tidak ditemukan")
            continue
        target = OUTPUT_DIR / f"{source.stem}.docx"
        convert(source, target)
        print(f"[ok] {source.name} -> {target.name}")
        converted += 1
    print(f"Selesai: {converted} dokumen di {OUTPUT_DIR}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
